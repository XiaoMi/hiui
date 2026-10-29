#!/usr/bin/env python3
"""
UX Walkthrough - 本地 DOCX 生成标准入口

用途：
- 读取结构化 report.json
- 统一生成本地 .docx 报告
- 返回结构化 JSON 结果，供外层决定 success / failed 收口
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
except ImportError as exc:  # pragma: no cover - import guard
    Document = None  # type: ignore[assignment]
    WD_ALIGN_PARAGRAPH = None  # type: ignore[assignment]
    Inches = None  # type: ignore[assignment]
    Pt = None  # type: ignore[assignment]
    DOCX_IMPORT_ERROR = exc
else:
    DOCX_IMPORT_ERROR = None


SUPPORTED_SOURCES = {"code", "url", "screenshot"}
MAX_IMAGE_WIDTH_INCHES = 6.2
BUSINESS_SOURCE_LABELS = {
    "inferred": "页面 / 代码 / 截图推断",
    "mixed": "文档与现状核对",
    "prd": "PRD / 需求文档",
    "unknown": "信息不足",
}
BUSINESS_CONFIDENCE_LABELS = {"high": "高", "low": "低", "medium": "中"}
FLOW_STATUS_LABELS = {"confirmed": "已确认", "inferred": "推断", "pending": "待确认"}
ISSUE_CONFIDENCE_LABELS = {"confirmed": "已确认", "inferred": "推断", "pending": "待确认"}
CHECKLIST_HEADER_RE = re.compile(r"^### (UX-\d+\.\d+)\s+(.+)$", re.MULTILINE)
CHECKLIST_STATUS_LABELS = {
    "issue": "命中问题",
    "n/a": "不适用",
    "pass": "通过",
    "pending": "待交互验证",
    "unknown": "未识别",
}
STYLE_STATUS_LABELS = CHECKLIST_STATUS_LABELS.copy()
HIUI_ALLOWED_SEVERITIES = {"P0", "P1", "P2", "P3"}


def _skill_root() -> Path:
    return Path(__file__).resolve().parent.parent


def _build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="根据结构化 report.json 生成 ux-walkthrough 本地 docx 报告")
    parser.add_argument("--report-json", required=True)
    parser.add_argument("--output", required=True)
    parser.add_argument("--overwrite", action="store_true")
    parser.add_argument(
        "--skip-annotation-check",
        action="store_true",
        help="已禁用；标注门禁不得跳过",
    )
    parser.add_argument(
        "--skip-checklist-check",
        action="store_true",
        help="已禁用；checklist 覆盖门禁不得跳过",
    )
    parser.add_argument("--json", action="store_true")
    return parser


def _fail(message: str, detail: str = "", *, reason: str = "tool_error", extra: dict | None = None) -> dict:
    payload = {
        "detail": detail[:500],
        "failure_reason": reason,
        "failure_stage": "docx_generation",
        "message": message,
        "ok": False,
        "status": "failed",
    }
    if extra:
        payload.update(extra)
    return payload


def _ok(message: str, *, output_path: str, file_name: str, source: str, issue_count: int, image_count: int) -> dict:
    return {
        "file_name": file_name,
        "image_count": image_count,
        "issue_count": issue_count,
        "message": message,
        "ok": True,
        "output_path": output_path,
        "source": source,
        "status": "generated",
    }


def _load_json(path: Path) -> dict:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except FileNotFoundError as exc:
        raise ValueError(f"未找到 report.json：{path}") from exc
    except json.JSONDecodeError as exc:
        raise ValueError(f"report.json 不是合法 JSON：{exc}") from exc
    except OSError as exc:
        raise ValueError(f"读取 report.json 失败：{exc}") from exc

    if not isinstance(data, dict):
        raise ValueError("report.json 根节点必须是对象")
    return data


def _text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _string_list(value: object) -> list[str]:
    if isinstance(value, list):
        result = []
        for item in value:
            text = _text(item)
            if text:
                result.append(text)
        return result
    text = _text(value)
    return [text] if text else []


def _format_list(items: list[str], sep: str = "、") -> str:
    return sep.join(item for item in items if item)


def _normalize_source(raw: object) -> str:
    source = _text(raw).lower()
    if source not in SUPPORTED_SOURCES:
        raise ValueError("report.json 中的 source 必须是 code / url / screenshot 之一")
    return source


def _resolve_output_path(raw_output: str) -> Path:
    output = Path(raw_output).expanduser()
    if output.suffix.lower() != ".docx":
        raise ValueError("--output 必须以 .docx 结尾")
    return output.resolve()


def _load_checklist_items(checklist_path: Path | None = None) -> list[tuple[str, str]]:
    path = checklist_path or (_skill_root() / "references" / "ux-checklist.md")
    text = path.read_text(encoding="utf-8")
    items: list[tuple[str, str]] = []
    seen: set[str] = set()
    for item_id, item_name in CHECKLIST_HEADER_RE.findall(text):
        if item_id in seen:
            continue
        items.append((item_id, _text(item_name)))
        seen.add(item_id)
    if not items:
        raise ValueError(f"未能从 {path} 解析 checklist 标题")
    return items


def _parse_checklist_status(raw: object) -> tuple[str, str]:
    text = _text(raw)
    if not text:
        return "", ""

    lowered = text.lower()
    if lowered == "issue":
        return "issue", "本项已命中体验问题，详见正文问题详情。"
    if lowered == "pass":
        return "pass", "本项已检查，未发现需要单列的问题。"
    if lowered.startswith("pass:"):
        return "pass", text.split(":", 1)[1].strip()
    if lowered.startswith("pending:"):
        return "pending", text.split(":", 1)[1].strip()
    if lowered.startswith("n/a:") or lowered.startswith("na:"):
        return "n/a", text.split(":", 1)[1].strip()
    return "unknown", text


def _normalize_checklist_coverage(data: dict) -> list[dict]:
    raw = data.get("checklist_coverage")
    if not isinstance(raw, dict):
        return []

    try:
        checklist_items = _load_checklist_items()
    except (OSError, ValueError):
        checklist_items = []
        for key in sorted(raw.keys(), key=lambda item: _text(item)):
            item_id = _text(key)
            if item_id.startswith("UX-"):
                checklist_items.append((item_id, item_id))

    rows: list[dict] = []
    for item_id, item_name in checklist_items:
        status, note = _parse_checklist_status(raw.get(item_id, ""))
        if not status:
            status = "unknown"
            note = "本项在报告中未写明结果。"
        rows.append(
            {
                "id": item_id,
                "name": item_name,
                "note": note,
                "status": status,
                "status_label": CHECKLIST_STATUS_LABELS.get(status, CHECKLIST_STATUS_LABELS["unknown"]),
            }
        )
    return rows


def _resolve_image_path(raw: str, base_dir: Path) -> Path:
    path = Path(raw).expanduser()
    if not path.is_absolute():
        path = (base_dir / path).resolve()
    return path


def _normalize_image_entry(entry: object, base_dir: Path) -> dict:
    if isinstance(entry, str):
        path_text = _text(entry)
        caption = ""
    elif isinstance(entry, dict):
        path_text = _text(entry.get("path", ""))
        caption = _text(entry.get("caption", ""))
    else:
        raise ValueError("images 中的每一项必须是字符串路径或包含 path 的对象")

    if not path_text:
        raise ValueError("images 中存在空路径")

    path = _resolve_image_path(path_text, base_dir)
    if not path.is_file():
        raise ValueError(f"图片不存在：{path}")

    return {
        "caption": caption,
        "path": path,
    }


def _normalize_issue(item: object, index: int, base_dir: Path) -> dict:
    if not isinstance(item, dict):
        raise ValueError(f"issues[{index}] 必须是对象")

    title = _text(item.get("title", ""))
    if not title:
        raise ValueError(f"issues[{index}] 缺少 title")

    severity = _text(item.get("severity", "")).upper() or "P2"
    category = _text(item.get("category", ""))
    allowed_severities = {"P0", "P1", "P2"}
    if category == "样式问题":
        allowed_severities.add("P3")
    if severity not in allowed_severities:
        allowed_text = "P0 / P1 / P2 / P3" if category == "样式问题" else "P0 / P1 / P2"
        raise ValueError(f"issues[{index}] 的 severity 必须是 {allowed_text}")

    description = _text(item.get("description", "")) or _text(item.get("problem_description", ""))
    if not description:
        raise ValueError(f"issues[{index}] 缺少 description")

    suggestion = _text(item.get("suggestion", "")) or _text(item.get("improvement", ""))
    location = _text(item.get("location", "")) or _text(item.get("position", ""))
    flow_step = _text(item.get("flow_step", "")) or _text(item.get("journey_step", ""))
    affected_users = _string_list(item.get("affected_users", [])) or _string_list(item.get("affected_user", []))
    business_impact = _text(item.get("business_impact", ""))
    confidence = _text(item.get("confidence", "")).lower()
    if confidence not in ISSUE_CONFIDENCE_LABELS:
        confidence = ""
    evidence_note = _text(item.get("evidence_note", "")) or _text(item.get("evidence", ""))
    raw_images = item.get("images", [])
    if isinstance(raw_images, (str, dict)):
        raw_images = [raw_images]
    elif not isinstance(raw_images, list):
        raise ValueError(f"issues[{index}] 的 images 必须是数组、字符串路径或包含 path 的对象")
    images = [_normalize_image_entry(entry, base_dir) for entry in raw_images]
    item_index = int(item.get("index", index))

    return {
        "affected_users": affected_users,
        "business_impact": business_impact,
        "category": category,
        "confidence": confidence,
        "description": description,
        "evidence_note": evidence_note,
        "flow_step": flow_step,
        "images": images,
        "index": item_index,
        "location": location,
        "severity": severity,
        "suggestion": suggestion,
        "title": title,
    }


def _normalize_business_context(data: dict) -> dict:
    raw = data.get("business_context", {})
    if not isinstance(raw, dict):
        raw = {}

    source = _text(raw.get("source", "")).lower()
    if source not in BUSINESS_SOURCE_LABELS:
        source = ""

    confidence = _text(raw.get("confidence", "")).lower()
    if confidence not in BUSINESS_CONFIDENCE_LABELS:
        confidence = ""

    return {
        "assumptions": _string_list(raw.get("assumptions", [])),
        "confidence": confidence,
        "core_tasks": _string_list(raw.get("core_tasks", [])),
        "missing_information": _string_list(raw.get("missing_information", [])),
        "primary_users": _string_list(raw.get("primary_users", [])),
        "product_goal": _text(raw.get("product_goal", "")),
        "risk_points": _string_list(raw.get("risk_points", [])),
        "source": source,
        "success_signal": _text(raw.get("success_signal", "")),
    }


def _normalize_style_dimension(item: object, index: int) -> dict:
    if not isinstance(item, dict):
        raise ValueError(f"style_review.dimensions[{index}] 必须是对象")

    name = _text(item.get("name", "")) or _text(item.get("dimension", "")) or f"维度 {index}"
    status = _text(item.get("status", "")).lower()
    if status not in STYLE_STATUS_LABELS:
        status = "unknown"
    note = _text(item.get("note", "")) or _text(item.get("summary", ""))
    evidence = _text(item.get("evidence", "")) or _text(item.get("evidence_note", ""))
    return {
        "evidence": evidence,
        "name": name,
        "note": note,
        "status": status,
        "status_label": STYLE_STATUS_LABELS.get(status, STYLE_STATUS_LABELS["unknown"]),
    }


def _normalize_hiui_component_check(item: object, index: int) -> dict:
    if not isinstance(item, dict):
        raise ValueError(f"style_review.hiui_component_checks[{index}] 必须是对象")

    check_id = _text(item.get("id", "")) or f"HIUI-CHECK-{index}"
    name = _text(item.get("name", "")) or check_id
    status = _text(item.get("status", "")).lower()
    if status not in STYLE_STATUS_LABELS:
        status = "unknown"
    note = _text(item.get("note", "")) or _text(item.get("summary", ""))
    evidence = _text(item.get("evidence", "")) or _text(item.get("evidence_note", ""))
    return {
        "evidence": evidence,
        "id": check_id,
        "name": name,
        "note": note,
        "status": status,
        "status_label": STYLE_STATUS_LABELS.get(status, STYLE_STATUS_LABELS["unknown"]),
    }


def _normalize_hiui_component_checks(raw: object) -> list[dict]:
    if not raw:
        return []
    if isinstance(raw, dict):
        iterable = []
        for check_id, value in raw.items():
            item = dict(value) if isinstance(value, dict) else {"status": value}
            item.setdefault("id", check_id)
            iterable.append(item)
    elif isinstance(raw, list):
        iterable = raw
    else:
        raise ValueError("report.json 中的 style_review.hiui_component_checks 必须是数组或对象")
    return [_normalize_hiui_component_check(item, idx + 1) for idx, item in enumerate(iterable)]


def _normalize_style_review(data: dict) -> dict:
    raw = data.get("style_review", {})
    if not isinstance(raw, dict):
        raw = {}

    dimensions_raw = raw.get("dimensions", [])
    if dimensions_raw and not isinstance(dimensions_raw, list):
        raise ValueError("report.json 中的 style_review.dimensions 必须是数组")
    dimensions = [_normalize_style_dimension(item, idx + 1) for idx, item in enumerate(dimensions_raw or [])]

    hiui_sections_raw = raw.get("hiui_sections", [])
    if hiui_sections_raw and not isinstance(hiui_sections_raw, list):
        raise ValueError("report.json 中的 style_review.hiui_sections 必须是数组")
    hiui_sections = [_normalize_hiui_style_section(item, idx + 1) for idx, item in enumerate(hiui_sections_raw or [])]
    hiui_component_checks = _normalize_hiui_component_checks(raw.get("hiui_component_checks", []))

    return {
        "framework": _text(raw.get("framework", "")).lower(),
        "dimensions": dimensions,
        "hiui_sections": hiui_sections,
        "hiui_component_checks": hiui_component_checks,
        "priority_actions": _string_list(raw.get("priority_actions", [])),
        "summary": _text(raw.get("summary", "")),
    }


def _is_hiui_style_review(style_review: dict) -> bool:
    framework = _text(style_review.get("framework", "")).lower()
    return framework == "hiui-design"


def _normalize_hiui_style_section(item: object, index: int) -> dict:
    if not isinstance(item, dict):
        raise ValueError(f"style_review.hiui_sections[{index}] 必须是对象")

    name = _text(item.get("name", "")) or f"HiUI 对照项 {index}"
    status = _text(item.get("status", "")).lower()
    if status not in STYLE_STATUS_LABELS:
        status = "unknown"
    raw_severity = _text(item.get("severity", "")).upper()
    if raw_severity and raw_severity not in HIUI_ALLOWED_SEVERITIES:
        raise ValueError(
            f"style_review.hiui_sections[{index}] 的 severity 必须是 P0 / P1 / P2 / P3"
        )
    severity = raw_severity
    note = _text(item.get("note", "")) or _text(item.get("summary", ""))
    evidence = _text(item.get("evidence", "")) or _text(item.get("evidence_note", ""))
    return {
        "evidence": evidence,
        "name": name,
        "note": note,
        "severity": severity,
        "status": status,
        "status_label": STYLE_STATUS_LABELS.get(status, STYLE_STATUS_LABELS["unknown"]),
    }


def _normalize_core_flow(item: object, index: int) -> dict:
    if not isinstance(item, dict):
        raise ValueError(f"core_flows[{index}] 必须是对象")

    name = _text(item.get("name", "")) or f"流程 {index}"
    goal = _text(item.get("goal", ""))
    steps = _string_list(item.get("steps", []))
    status = _text(item.get("status", "")).lower()
    if status not in FLOW_STATUS_LABELS:
        status = ""

    return {
        "goal": goal,
        "name": name,
        "status": status,
        "steps": steps,
    }


def _normalize_report(data: dict, report_path: Path) -> dict:
    source = _normalize_source(data.get("source", ""))
    title = _text(data.get("title", "")) or f"{_text(data.get('project', 'UX Walkthrough')) or 'UX Walkthrough'} 体验走查报告"
    project = _text(data.get("project", ""))
    generated_at = _text(data.get("generated_at", ""))
    scope = _text(data.get("scope", "")) or _text(data.get("walkthrough_scope", ""))
    tech_stack = _text(data.get("tech_stack", ""))

    summary = data.get("summary", {}) if isinstance(data.get("summary"), dict) else {}
    overall = _text(summary.get("overall", "")) or _text(data.get("overall_evaluation", ""))
    key_points = _string_list(summary.get("key_points", [])) or _string_list(data.get("key_points", []))
    improvement_direction = _text(summary.get("improvement_direction", "")) or _text(data.get("improvement_direction", ""))
    business_context = _normalize_business_context(data)
    style_review = _normalize_style_review(data)
    core_flows_raw = data.get("core_flows", [])
    if core_flows_raw and not isinstance(core_flows_raw, list):
        raise ValueError("report.json 中的 core_flows 必须是数组")
    core_flows = [_normalize_core_flow(item, idx + 1) for idx, item in enumerate(core_flows_raw or [])]

    issues_raw = data.get("issues", [])
    if not isinstance(issues_raw, list) or not issues_raw:
        raise ValueError("report.json 中的 issues 必须是非空数组")
    issues = [_normalize_issue(item, idx + 1, report_path.parent) for idx, item in enumerate(issues_raw)]

    appendix = data.get("appendix", {}) if isinstance(data.get("appendix"), dict) else {}
    appendix_note = _text(appendix.get("note", "")) or _text(data.get("appendix_note", ""))
    appendix_items = _string_list(appendix.get("items", []))
    checklist_rows = _normalize_checklist_coverage(data)

    return {
        "appendix_items": appendix_items,
        "appendix_note": appendix_note,
        "business_context": business_context,
        "checklist_rows": checklist_rows,
        "core_flows": core_flows,
        "generated_at": generated_at,
        "improvement_direction": improvement_direction,
        "issues": issues,
        "key_points": key_points,
        "overall": overall,
        "project": project,
        "scope": scope,
        "source": source,
        "style_review": style_review,
        "tech_stack": tech_stack,
        "title": title,
    }


def _set_base_style(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    normal = document.styles["Normal"]
    normal.font.name = "PingFang SC"
    normal.font.size = Pt(10.5)


def _add_title(document: Document, report: dict) -> None:
    title = document.add_paragraph()
    title.style = document.styles["Title"]
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(report["title"])
    run.font.size = Pt(18)

    meta_lines = [
        f"走查时间：{report['generated_at'] or '未提供'}",
        f"走查模式：{report['source']}",
        f"走查范围：{report['scope'] or '未提供'}",
    ]
    if report["project"]:
        meta_lines.insert(0, f"项目：{report['project']}")
    if report["tech_stack"]:
        meta_lines.append(f"技术栈：{report['tech_stack']}")

    for line in meta_lines:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(line)


def _add_summary(document: Document, report: dict) -> None:
    document.add_heading("体验问题总结", level=1)
    if report["overall"]:
        document.add_heading("整体评价", level=2)
        document.add_paragraph(report["overall"])

    if report["key_points"]:
        document.add_heading("核心问题", level=2)
        for item in report["key_points"]:
            document.add_paragraph(item, style="List Bullet")

    if report["improvement_direction"]:
        document.add_heading("改进方向", level=2)
        document.add_paragraph(report["improvement_direction"])


def _add_style_review(document: Document, report: dict, *, heading_level: int = 1) -> None:
    style_review = report["style_review"]
    if not style_review["dimensions"] and not style_review["hiui_component_checks"]:
        return

    document.add_heading("样式专项观察", level=heading_level)
    if style_review["dimensions"]:
        table = document.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "维度"
        header[1].text = "结论"
        header[2].text = "证据"
        for item in style_review["dimensions"]:
            row = table.add_row().cells
            row[0].text = item["name"]
            row[1].text = f"{item['status_label']}：{item['note']}" if item["note"] else item["status_label"]
            row[2].text = item["evidence"]

    if style_review["hiui_component_checks"]:
        document.add_paragraph("HiUI 组件逐项检查")
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "检查项"
        header[1].text = "结论"
        header[2].text = "证据"
        header[3].text = "ID"
        for item in style_review["hiui_component_checks"]:
            row = table.add_row().cells
            row[0].text = item["name"]
            row[1].text = f"{item['status_label']}：{item['note']}" if item["note"] else item["status_label"]
            row[2].text = item["evidence"]
            row[3].text = item["id"]


def _build_users_tasks_summary(business_context: dict) -> str:
    parts: list[str] = []
    if business_context["primary_users"]:
        parts.append(f"用户：{_format_list(business_context['primary_users'])}")
    if business_context["core_tasks"]:
        parts.append(f"任务：{_format_list(business_context['core_tasks'], sep='；')}")
    return "；".join(parts)


def _build_flow_summary(core_flows: list[dict]) -> str:
    snippets: list[str] = []
    for flow in core_flows[:2]:
        steps = " → ".join(flow["steps"][:4])
        if flow["name"] and steps:
            snippets.append(f"{flow['name']}：{steps}")
        elif flow["name"]:
            snippets.append(flow["name"])
        elif steps:
            snippets.append(steps)
    return "；".join(snippets)


def _build_boundary_summary(business_context: dict) -> str:
    parts: list[str] = []
    if business_context["source"]:
        parts.append(
            f"业务理解基于{BUSINESS_SOURCE_LABELS.get(business_context['source'], business_context['source'])}"
        )
    if business_context["assumptions"]:
        parts.append(f"当前假设：{_format_list(business_context['assumptions'], sep='；')}")
    if business_context["missing_information"]:
        parts.append(f"仍缺：{_format_list(business_context['missing_information'], sep='；')}")
    return "；".join(parts)


def _add_business_context(document: Document, report: dict) -> None:
    business_context = report["business_context"]
    core_flows = report["core_flows"]
    has_context = any(
        [
            business_context["source"],
            business_context["product_goal"],
            business_context["primary_users"],
            business_context["core_tasks"],
            business_context["assumptions"],
            business_context["missing_information"],
            core_flows,
        ]
    )
    if not has_context:
        return

    document.add_heading("业务理解简述", level=1)
    _add_issue_field(document, "产品目标", business_context["product_goal"])
    _add_issue_field(document, "主要用户与任务", _build_users_tasks_summary(business_context))
    _add_issue_field(document, "核心流程", _build_flow_summary(core_flows))
    _add_issue_field(document, "判断边界", _build_boundary_summary(business_context))


def _issue_heading(issue: dict) -> str:
    title = issue["title"]
    prefix = f"{issue['index']}. [{issue['severity']}]"
    if title.startswith(prefix):
        return title
    return f"{prefix}{title}" if title.startswith(" ") else f"{prefix} {title}"


def _add_issue_field(document: Document, label: str, value: str) -> None:
    if not value:
        return
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}：").bold = True
    paragraph.add_run(value)


def _add_issue_images(document: Document, issue: dict) -> int:
    image_count = 0
    for image in issue["images"]:
        document.add_picture(str(image["path"]), width=Inches(MAX_IMAGE_WIDTH_INCHES))
        image_count += 1
        if image["caption"]:
            caption = document.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            caption.add_run(image["caption"]).italic = True
    return image_count


def _add_issue_detail(document: Document, issue: dict) -> int:
    document.add_heading(_issue_heading(issue), level=2)
    if issue["category"]:
        _add_issue_field(document, "问题类型", issue["category"])
    _add_issue_field(document, "位置", issue["location"])
    _add_issue_field(document, "问题描述", issue["description"])
    _add_issue_field(document, "改进建议", issue["suggestion"])
    _add_issue_field(document, "证据说明", issue["evidence_note"])
    return _add_issue_images(document, issue)


def _distribution_rows(issues: list[dict]) -> tuple[list[tuple[str, int]], list[tuple[str, int, int, int, int]]]:
    category_counts: dict[str, int] = {}
    severity_by_category: dict[str, dict[str, int]] = {}

    for issue in issues:
        category = issue["category"] or "未分类"
        category_counts[category] = category_counts.get(category, 0) + 1
        severity_bucket = severity_by_category.setdefault(category, {"P0": 0, "P1": 0, "P2": 0, "P3": 0})
        severity_bucket[issue["severity"]] += 1

    top_categories = sorted(category_counts.items(), key=lambda item: (-item[1], item[0]))[:3]
    severity_rows = [
        (category, bucket["P0"], bucket["P1"], bucket["P2"], bucket["P3"])
        for category, bucket in sorted(severity_by_category.items(), key=lambda item: item[0])
    ]
    return top_categories, severity_rows


def _add_distribution(document: Document, issues: list[dict]) -> None:
    top_categories, severity_rows = _distribution_rows(issues)
    document.add_heading("问题分布", level=1)

    if top_categories:
        document.add_heading("Top 3 高频问题类型", level=2)
        table = document.add_table(rows=1, cols=3)
        header = table.rows[0].cells
        header[0].text = "排名"
        header[1].text = "问题类型"
        header[2].text = "出现次数"
        for idx, (category, count) in enumerate(top_categories, start=1):
            row = table.add_row().cells
            row[0].text = str(idx)
            row[1].text = category
            row[2].text = str(count)

    if severity_rows:
        document.add_heading("按问题类型 × 优先级分布", level=2)
        table = document.add_table(rows=1, cols=6)
        header = table.rows[0].cells
        header[0].text = "问题类型"
        header[1].text = "P0"
        header[2].text = "P1"
        header[3].text = "P2"
        header[4].text = "P3"
        header[5].text = "合计"

        total_p0 = total_p1 = total_p2 = total_p3 = 0
        for category, p0, p1, p2, p3 in severity_rows:
            row = table.add_row().cells
            row[0].text = category
            row[1].text = str(p0)
            row[2].text = str(p1)
            row[3].text = str(p2)
            row[4].text = str(p3)
            row[5].text = str(p0 + p1 + p2 + p3)
            total_p0 += p0
            total_p1 += p1
            total_p2 += p2
            total_p3 += p3

        total_row = table.add_row().cells
        total_row[0].text = "合计"
        total_row[1].text = str(total_p0)
        total_row[2].text = str(total_p1)
        total_row[3].text = str(total_p2)
        total_row[4].text = str(total_p3)
        total_row[5].text = str(total_p0 + total_p1 + total_p2 + total_p3)


def _add_appendix(document: Document, report: dict, total_images: int) -> None:
    document.add_heading("附录", level=1)
    document.add_paragraph(f"走查模式：{report['source']}")
    if report["project"]:
        document.add_paragraph(f"项目：{report['project']}")
    if report["scope"]:
        document.add_paragraph(f"走查范围：{report['scope']}")
    if report["appendix_note"]:
        document.add_paragraph(report["appendix_note"])
    for item in report["appendix_items"]:
        document.add_paragraph(item, style="List Bullet")
    if report["checklist_rows"]:
        document.add_heading("Checklist 逐条检查记录", level=2)
        table = document.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        header = table.rows[0].cells
        header[0].text = "ID"
        header[1].text = "检查项"
        header[2].text = "结果"
        header[3].text = "说明"
        for item in report["checklist_rows"]:
            row = table.add_row().cells
            row[0].text = item["id"]
            row[1].text = item["name"]
            row[2].text = item["status_label"]
            row[3].text = item["note"]
    _add_style_review(document, report, heading_level=2)
    if total_images == 0:
        document.add_paragraph("本次报告无截图证据。")


def _generate_docx(report: dict, output_path: Path) -> dict:
    if Document is None or WD_ALIGN_PARAGRAPH is None or Inches is None or Pt is None:
        return _fail(
            "DOCX generation failed: python-docx not available",
            detail=f"缺少依赖：{DOCX_IMPORT_ERROR}",
        )

    document = Document()
    _set_base_style(document)
    _add_title(document, report)
    _add_business_context(document, report)
    _add_summary(document, report)
    document.add_heading("问题详情", level=1)

    total_images = 0
    for issue in report["issues"]:
        total_images += _add_issue_detail(document, issue)

    _add_distribution(document, report["issues"])
    _add_appendix(document, report, total_images)

    try:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(str(output_path))
    except OSError as exc:
        return _fail("DOCX generation failed", detail=f"写入 docx 文件失败：{exc}")
    except Exception as exc:  # pragma: no cover - library specific
        return _fail("DOCX generation failed", detail=f"docx 库执行失败：{exc}")

    return _ok(
        "DOCX generated",
        file_name=output_path.name,
        image_count=total_images,
        issue_count=len(report["issues"]),
        output_path=str(output_path),
        source=report["source"],
    )


def build_result(args: argparse.Namespace) -> dict:
    if args.skip_annotation_check:
        return _fail(
            "DOCX generation failed: annotation gate cannot be skipped",
            detail="--skip-annotation-check 已禁用；正式 docx 生成必须通过标注门禁",
            reason="annotation_gate_skip_forbidden",
        )

    if args.skip_checklist_check:
        return _fail(
            "DOCX generation failed: checklist gate cannot be skipped",
            detail="--skip-checklist-check 已禁用；正式 docx 生成必须通过 checklist 覆盖门禁",
            reason="checklist_gate_skip_forbidden",
        )

    output_path = _resolve_output_path(args.output)
    if output_path.exists() and not args.overwrite:
        return _fail(
            "DOCX generation failed: output exists",
            detail=f"输出文件已存在：{output_path}；若需覆盖，请显式传 --overwrite",
            reason="tool_error",
            extra={"output_path": str(output_path)},
        )

    report_path = Path(args.report_json).expanduser().resolve()
    try:
        raw_report = _load_json(report_path)
        report = _normalize_report(raw_report, report_path)
    except ValueError as exc:
        return _fail("DOCX generation failed: invalid report json", detail=str(exc))

    try:
        from validate_core_flows import validate_report as validate_core_flows
    except ImportError as exc:
        return _fail(
            "DOCX generation failed: core-flow validator missing",
            detail=str(exc),
            reason="core_flow_gate_missing",
        )
    core_flow_gate = validate_core_flows(raw_report)
    if not core_flow_gate.get("ok"):
        errors = core_flow_gate.get("errors", [])
        detail = "；".join(errors[:5])
        if len(errors) > 5:
            detail += f"；等共 {len(errors)} 项"
        return _fail(
            "DOCX generation failed: core-flow gate not passed",
            detail=detail or "核心流程执行门禁未通过",
            reason="core_flow_gate_incomplete",
            extra={"core_flow_errors": errors},
        )

    try:
        from validate_style_review import validate_report as validate_style_review
    except ImportError as exc:
        return _fail(
            "DOCX generation failed: style-review validator missing",
            detail=str(exc),
            reason="style_review_gate_missing",
        )
    style_gate = validate_style_review(raw_report)
    if not style_gate.get("ok"):
        errors = style_gate.get("errors", [])
        detail = "；".join(errors[:5])
        if len(errors) > 5:
            detail += f"；等共 {len(errors)} 项"
        return _fail(
            "DOCX generation failed: style-review gate not passed",
            detail=detail or "样式走查门禁未通过",
            reason="style_review_incomplete",
            extra={"style_review_errors": errors},
        )

    try:
        from validate_checklist_coverage import validate_report as validate_checklist_coverage
    except ImportError as exc:
        return _fail(
            "DOCX generation failed: checklist validator missing",
            detail=str(exc),
        )
    coverage_gate = validate_checklist_coverage(raw_report)
    if not coverage_gate.get("ok"):
        errors = coverage_gate.get("errors", [])
        detail = "；".join(errors[:5])
        if len(errors) > 5:
            detail += f"；等共 {len(errors)} 项"
        return _fail(
            "DOCX generation failed: checklist coverage gate not passed",
            detail=detail or "checklist 覆盖门禁未通过",
            reason="checklist_coverage_incomplete",
            extra={"checklist_coverage_errors": errors},
        )

    try:
        from validate_report_annotations import validate_report
    except ImportError as exc:
        return _fail(
            "DOCX generation failed: annotation validator missing",
            detail=str(exc),
        )
    gate = validate_report(raw_report)
    if not gate.get("ok"):
        errors = gate.get("errors", [])
        detail = "；".join(errors[:5])
        if len(errors) > 5:
            detail += f"；等共 {len(errors)} 项"
        return _fail(
            "DOCX generation failed: annotation gate not passed",
            detail=detail or "标注门禁未通过",
            reason="annotations_missing",
            extra={"annotation_errors": errors},
        )

    return _generate_docx(report, output_path)


def main() -> int:
    args = _build_parser().parse_args()
    result = build_result(args)
    if args.json or not result["ok"]:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(result["output_path"])
    return 0 if result["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
